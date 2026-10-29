"""
Build samples/Carrier_Contacts.xlsx: the contact list the flow reads and the log it writes.

Both are real Excel tables (Insert > Table), because the Excel Online (Business)
connector can only read and write tables. All carriers and addresses are
fictional; the .example domain is reserved and never delivers mail.

Carrier IDs match the files produced by invoice-backup-splitter-vba. NSTPAR is
left out on purpose, so a test run shows the "contact not found" branch.

    python tools/build_samples.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parents[1] / "samples" / "Carrier_Contacts.xlsx"

CARRIERS = [
    # Carrier ID, Carrier Name, To, CC, Language
    ("LEVEXP", "LEVANTE EXPRESS SL", "facturacion@levante-express.example", "", "ES"),
    ("MESHAU", "MESETA HAULAGE SA", "billing@meseta-haulage.example", "ops@meseta-haulage.example", "ES"),
    ("NORCAR", "NORTE CARGO GMBH", "rechnungen@norte-cargo.example", "", "EN"),
    ("COSFRT", "COSTA FREIGHT BV", "invoices@costa-freight.example", "", "EN"),
    ("RIOTRP", "RIO TRANSPORTES LDA", "faturas@rio-transportes.example", "", "EN"),
    ("BALRAIL", "BALTIC RAIL LOGISTICS AB", "ap@baltic-rail.example", "finance@baltic-rail.example", "EN"),
    ("NSTFRT", "NORTHSTAR FREIGHT GMBH", "billing@northstar-freight.example", "", "EN"),
]


def add_table(ws, name, headers, rows):
    ws.append(headers)
    for r in rows or [[None] * len(headers)]:
        ws.append(list(r))
    ref = f"A1:{chr(64 + len(headers))}{ws.max_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[chr(64 + i)].width = max(14, len(h) + 4)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Carriers"
    add_table(ws, "Carriers", ["Carrier ID", "Carrier Name", "To", "CC", "Language"], CARRIERS)
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 32

    log = wb.create_sheet("SendLog")
    add_table(log, "SendLog", ["Date", "Carrier ID", "Files", "Sent To", "Result"], None)
    log.column_dimensions["C"].width = 80
    log.column_dimensions["D"].width = 38

    notes = wb.create_sheet("README")
    for i, line in enumerate([
        "Contact list and send log for the Carrier Invoice Mailer flow.",
        "Carriers: one row per carrier. Carrier ID must match the 2nd word of the export file name.",
        "Language: ES or EN, picks the email template. CC is optional.",
        "SendLog: written by the flow, one row per carrier emailed (SENT or TEST) and per file with no contact (CONTACT NOT FOUND).",
        "All carriers and addresses are fictional (.example domains never deliver mail).",
    ], start=1):
        notes.cell(row=i, column=1, value=line).font = Font(bold=(i == 1))
    notes.column_dimensions["A"].width = 100

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
